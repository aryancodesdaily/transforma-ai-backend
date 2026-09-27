import re
from pathlib import Path

from pypdf import PdfReader

from gemini_engine import extract_content_from_image


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".doc",
    ".docx",
    ".txt",
    ".pptx",
    ".md",
    ".csv",
    ".rtf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}

IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


def process_text(content: str) -> str:
    """
    Validate and clean plain text input.
    """

    if not content or not content.strip():
        raise ValueError("Source content cannot be empty.")

    return content.strip()


def process_file(file_path: str) -> str:
    """
    Extract textual content from a supported file.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        return process_pdf(path)

    if extension == ".docx":
        return process_docx(path)

    if extension in {".txt", ".md", ".csv", ".rtf"}:
        return process_plain_text_file(path)

    if extension in IMAGE_MIME_TYPES:
        return process_image(path)

    if extension == ".doc":
        raise NotImplementedError(
            "Legacy .doc files are not supported yet. "
            "Please use .docx format."
        )
    
    if extension == ".pptx":
        return process_pptx(path)

    raise ValueError(f"Unsupported file type: {extension}")


from pptx import Presentation


def process_pptx(file_path: str) -> str:
    """
    Extract all text (titles, body text, notes) from an uploaded .pptx
    and return it as a single plain-text string for the LLM pipeline.
    """

    prs = Presentation(file_path)
    parts = []

    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"Slide {i}:")

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        parts.append(text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")

    extracted = "\n".join(parts).strip()

    if not extracted:
        raise ValueError("No readable text found in the uploaded PPTX file.")

    return extracted




def process_pdf(path: Path) -> str:
    """
    Extract text from a PDF file.
    """

    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:
        try:
            text = page.extract_text(extraction_mode="layout")
        except Exception:
            text = page.extract_text()

        if text:
            text = re.sub(r"[\x7f\x80-\x9f]", "-", text)
            cleaned_lines = [line.rstrip() for line in text.splitlines()]
            cleaned_text = "\n".join(cleaned_lines).strip()
            if cleaned_text:
                pages.append(cleaned_text)

    content = "\n\n".join(pages)

    if not content.strip():
        raise ValueError(
            f"No readable text found in PDF: {path.name}"
        )

    return content.strip()

def process_docx(path: Path) -> str:
    """
    Extract text from a DOCX file.
    """

    from docx import Document

    document = Document(str(path))

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    content = "\n\n".join(paragraphs)

    if not content.strip():
        raise ValueError(
            f"No readable text found in DOCX: {path.name}"
        )

    return content.strip()

def process_plain_text_file(path: Path) -> str:
    """
    Read text-based files such as TXT, MD, CSV and RTF.
    """

    content = path.read_text(
        encoding="utf-8",
        errors="replace"
    )

    if not content.strip():
        raise ValueError(
            f"File is empty: {path.name}"
        )

    return content.strip()


def process_image(path: Path) -> str:
    """
    Extract/describe content from an image using Gemini vision, so it
    can be merged into the same source-content pipeline as PDFs, DOCX
    files, and plain text.
    """

    image_bytes = path.read_bytes()
    mime_type = IMAGE_MIME_TYPES[path.suffix.lower()]

    content = extract_content_from_image(image_bytes, mime_type)

    if not content.strip():
        raise ValueError(
            f"No readable content found in image: {path.name}"
        )

    return content.strip()


def process_multiple_files(file_paths: list[str]) -> str:
    """
    Extract and combine content from multiple files.
    """

    if not file_paths:
        raise ValueError("No files were provided.")

    processed_files = []

    for file_path in file_paths:
        content = process_file(file_path)

        processed_files.append(
            f"===== {Path(file_path).name} =====\n\n{content}"
        )

    return "\n\n".join(processed_files)